"""Общие фикстуры тестов."""

from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches

from project_doc_agent.config import ProjectConfig, make_project_config, save_project_config
from project_doc_agent.workspace import Workspace

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2nXQAAAAASUVORK5CYII="
)


def write_text(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")
    return path


def make_workspace(root: Path, source_dir: Path, profile: str = "integration") -> tuple[Workspace, ProjectConfig]:
    config = make_project_config(
        client="marta",
        client_name="ООО «Марта»",
        project="toplog-rc",
        project_name="Интеграция TopLog WMS — РЦ",
        bases=["ut11-prod"],
        profile=profile,
        provider="deepseek",
        source_dir=str(source_dir).replace("\\", "/"),
    )
    root.mkdir(parents=True, exist_ok=True)
    workspace = Workspace(root)
    workspace.configure(config.output_dir, config.reports_dir)
    workspace.ensure_layout()
    save_project_config(root, config)
    return workspace, config


@pytest.fixture()
def sample_source(tmp_path: Path) -> Path:
    source = tmp_path / "sources"
    source.mkdir()
    write_text(
        source / "alpha.md",
        """---
title: "Интеграционный дизайн"
author: "Тест"
---

# Интеграция TopLog

## Передача заказа

Метод createShipmentOrder передаёт заказ на отгрузку из КИС в WMS.
Номенклатура передаётся отдельным справочником.

| Поле | Тип |
| --- | --- |
| OrderID | string |
| Rows | array |

## Авторизация

Для обмена используется отдельный пользователь. Авторизация выполняется по логину и паролю.
""",
    )
    write_text(
        source / "beta.md",
        """# Интеграция TopLog

## Передача заказа

В этом документе указано, что createShipmentOrder вызывается со стороны WMS → КИС.
""",
    )
    write_text(
        source / "notes.txt",
        """Общие заметки по проекту.

Ошибки обрабатываются повторной выгрузкой пакета. Мониторинг статусов ведётся вручную.
""",
    )
    return source


@pytest.fixture()
def workspace(sample_source: Path, tmp_path: Path) -> tuple[Workspace, ProjectConfig]:
    return make_workspace(tmp_path / "project", sample_source)


@pytest.fixture()
def docx_source(tmp_path: Path) -> Path:
    source = tmp_path / "docx_sources"
    source.mkdir()
    document = Document()
    document.add_heading("Интеграция", level=1)
    document.add_heading("Передача заказа", level=2)
    document.add_paragraph("Метод createShipmentOrder передаёт заказ из КИС в WMS.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Поле"
    table.cell(0, 1).text = "Тип"
    table.cell(1, 0).text = "OrderID"
    table.cell(1, 1).text = "string"
    document.add_picture(io.BytesIO(PNG_1X1), width=Inches(0.2))
    document.add_heading("Ошибки", level=2)
    document.add_paragraph("Ошибки обрабатываются повторной отправкой.")
    document.save(source / "design.docx")
    return source
