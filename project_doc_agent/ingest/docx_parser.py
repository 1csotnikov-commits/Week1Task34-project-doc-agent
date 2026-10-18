"""Парсер DOCX: порядок абзацев/таблиц, TOC, outlineLvl, таблицы, изображения."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Any, Literal

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from project_doc_agent.ingest.common import BlockBuilder, ParseOutput
from project_doc_agent.models import ManualReviewItem
from project_doc_agent.utils import normalize_title_key, normalize_ws

HEADING_STYLE_RE = re.compile(r"^(?:heading|заголовок)\s*([1-9])$", re.IGNORECASE)
TITLE_STYLES = {"title", "title-major", "название"}
NUMBERED_RE = re.compile(r"^\d+(?:\.\d+){0,6}\.?\s+\S")
TOC_TITLE_RE = re.compile(r"^(оглавление|содержание)$", re.IGNORECASE)


@dataclass
class ParagraphInfo:
    element: Any
    text: str
    style: str
    outline_level: int | None
    bold: bool
    has_image: bool
    has_object: bool
    is_toc: bool
    heading_level: int | None = None
    heading_text: str = ""


@dataclass
class TocEntry:
    number: str
    title: str
    level: int


@dataclass
class DocxDocument:
    paragraphs: list[ParagraphInfo] = field(default_factory=list)
    toc: list[TocEntry] = field(default_factory=list)
    title: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


def _element_text(element: Any) -> str:
    parts: list[str] = []
    for node in element.iter():
        tag = node.tag
        if tag == qn("w:t"):
            parts.append(node.text or "")
        elif tag == qn("w:tab"):
            parts.append("\t")
        elif tag in {qn("w:br"), qn("w:cr")}:
            parts.append("\n")
    return "".join(parts)


def _paragraph_style(paragraph: Paragraph) -> str:
    try:
        return paragraph.style.name if paragraph.style is not None else ""
    except (AttributeError, KeyError):
        return ""


def _outline_level(element: Any) -> int | None:
    p_pr = getattr(element, "pPr", None)
    if p_pr is None:
        return None
    outline = getattr(p_pr, "outlineLvl", None)
    if outline is None:
        return None
    try:
        return int(outline.val)
    except (AttributeError, TypeError, ValueError):
        return None


def _all_bold(paragraph: Paragraph) -> bool:
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.bold is True for run in runs)


def _has_drawing(element: Any) -> bool:
    return bool(element.findall(".//" + qn("w:drawing")) or element.findall(".//" + qn("w:pict")))


def _has_object(element: Any) -> bool:
    return bool(element.findall(".//" + qn("w:object")) or element.findall(".//" + qn("w:oleObject")))


def _parse_toc_entry(text: str, style: str) -> TocEntry | None:
    parts = [part.strip() for part in text.replace("\n", "\t").split("\t") if part.strip()]
    if not parts:
        return None
    if len(parts) > 1 and parts[-1].isdigit():
        parts.pop()
    number = ""
    if len(parts) > 1 and re.fullmatch(r"\d+(?:\.\d+)*", parts[0]):
        number = parts.pop(0)
    title = normalize_ws(" ".join(parts))
    if not title:
        return None
    match = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.+)$", title)
    if match:
        number = number or match.group(1)
        title = normalize_ws(match.group(2))
    level_match = re.search(r"(\d+)", style)
    level = int(level_match.group(1)) if level_match else max(1, number.count(".") + 1 if number else 1)
    return TocEntry(number=number, title=title, level=level)


def _find_toc_entry(text: str, toc: list[TocEntry], cursor: int) -> int:
    key = normalize_title_key(text)
    if not key:
        return -1
    best_index = -1
    best_ratio = 0.0
    for index in range(cursor, min(len(toc), cursor + 12)):
        ratio = SequenceMatcher(None, key, normalize_title_key(toc[index].title)).ratio()
        if ratio > best_ratio:
            best_ratio = ratio
            best_index = index
    return best_index if best_ratio >= 0.8 else -1


def _collect_paragraphs(document: Any) -> DocxDocument:
    result = DocxDocument()
    for child in document.element.body.iterchildren():
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document)
        text = _element_text(child)
        style = _paragraph_style(paragraph)
        style_key = style.casefold()
        is_toc = style_key.startswith("toc") or style_key.startswith("оглавление")
        outline = _outline_level(child)
        info = ParagraphInfo(
            element=child,
            text=normalize_ws(text.replace("\t", " ")),
            style=style,
            outline_level=outline,
            bold=_all_bold(paragraph),
            has_image=_has_drawing(child),
            has_object=_has_object(child),
            is_toc=is_toc,
        )
        result.paragraphs.append(info)
        if is_toc:
            entry = _parse_toc_entry(text, style)
            if entry is not None:
                result.toc.append(entry)

    core = document.core_properties
    title = (core.title or "").strip() or None
    if title is None:
        title_candidates = [
            info.text for info in result.paragraphs if info.text and info.style.casefold() in TITLE_STYLES
        ]
        if title_candidates:
            title = max(title_candidates, key=len)
    if title is None:
        for info in result.paragraphs:
            if info.text and info.bold and len(info.text) <= 120 and not NUMBERED_RE.match(info.text):
                title = info.text
                break
    result.title = title
    result.metadata = {
        "author": core.author or "",
        "created": core.created.isoformat() if core.created else "",
        "modified": core.modified.isoformat() if core.modified else "",
        "subject": core.subject or "",
        "keywords": core.keywords or "",
        "category": core.category or "",
        "has_toc": bool(result.toc),
    }
    return result


def _explicit_headings(paragraphs: list[ParagraphInfo]) -> list[ParagraphInfo]:
    headings: list[ParagraphInfo] = []
    for info in paragraphs:
        if info.is_toc or not info.text or TOC_TITLE_RE.match(info.text):
            continue
        style_key = info.style.casefold()
        style_level = HEADING_STYLE_RE.match(info.style)
        if info.outline_level is not None:
            info.heading_level = info.outline_level + 1
        elif style_level:
            info.heading_level = int(style_level.group(1))
        elif style_key in {"heading", "заголовок"}:
            info.heading_level = 1
        else:
            continue
        headings.append(info)
    return headings


def _fallback_headings(paragraphs: list[ParagraphInfo], title: str | None) -> list[ParagraphInfo]:
    headings: list[ParagraphInfo] = []
    first_text_seen = False
    for info in paragraphs:
        if info.is_toc or not info.text:
            continue
        is_first = not first_text_seen
        first_text_seen = True
        if is_first and (title is None or info.text == title) and not NUMBERED_RE.match(info.text):
            continue
        if not info.bold or len(info.text) > 150:
            continue
        if info.text.rstrip().endswith((":", ";", ",")):
            continue
        match = re.match(r"^(\d+(?:\.\d+){0,5})\.?\s+\S", info.text)
        if not match:
            continue
        info.heading_level = min(6, len(match.group(1).split(".")))
        headings.append(info)
    return headings


def _assign_headings(data: DocxDocument) -> list[ParagraphInfo]:
    explicit = _explicit_headings(data.paragraphs)
    if explicit:
        cursor = 0
        for info in explicit:
            display = info.text
            toc_index = _find_toc_entry(info.text, data.toc, cursor)
            if toc_index >= 0:
                entry = data.toc[toc_index]
                cursor = toc_index + 1
                if entry.number and not NUMBERED_RE.match(info.text):
                    display = f"{entry.number} {info.text}"
                if info.outline_level is None and not HEADING_STYLE_RE.match(info.style):
                    info.heading_level = entry.level
            info.heading_text = display
        return explicit
    fallback = _fallback_headings(data.paragraphs, data.title)
    for info in fallback:
        info.heading_text = info.text
    return fallback


def _table_markdown(table: Table) -> tuple[str, int, int]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [normalize_ws(cell.text).replace("|", "\\|")[:1000] for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return "", 0, 0
    width = max(len(row) for row in rows)
    for row in rows:
        row.extend([""] * (width - len(row)))
    header = [(cell or f"Колонка {index + 1}") for index, cell in enumerate(rows[0])]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines), len(rows), width


def _paragraph_has_numbering(element: Any) -> bool:
    p_pr = getattr(element, "pPr", None)
    return bool(p_pr is not None and getattr(p_pr, "numPr", None) is not None)



def parse_docx(path: Any, document_id: str, file: str) -> ParseOutput:
    output = ParseOutput()
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 - повреждённый DOCX
        output.warnings.append(f"Не удалось прочитать DOCX {file}: {exc}")
        output.manual_review.append(
            ManualReviewItem(file=file, reason=f"DOCX не читается: {exc}", kind="unreadable")
        )
        return output
    data = _collect_paragraphs(document)
    output.metadata = dict(data.metadata)
    output.metadata["toc_entries"] = len(data.toc)
    output.title = data.title or file.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    headings = _assign_headings(data)
    output.metadata["headings"] = len(headings)
    if not headings:
        output.warnings.append(
            f"В {file} не найдены заголовки (стили, outlineLvl или нумерация): "
            "структура карточек будет обобщённой."
        )
    heading_map = {id(info.element): info for info in headings}
    paragraph_map = {id(info.element): info for info in data.paragraphs}
    builder = BlockBuilder(document_id, file)
    stack: list[tuple[int, str]] = []

    def current_path() -> list[str]:
        return [title for _, title in stack]

    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            info = paragraph_map.get(id(child))
            if info is None or info.is_toc:
                continue
            heading = heading_map.get(id(child))
            if heading is not None:
                level = heading.heading_level or 1
                stack = [item for item in stack if item[0] < level]
                stack.append((level, heading.heading_text or heading.text))
                builder.add("heading", heading.heading_text or heading.text, current_path(), level=level)
                continue
            if info.text and TOC_TITLE_RE.match(info.text):
                continue
            if info.text:
                style_key = info.style.casefold()
                kind: Literal["paragraph", "list"] = (
                    "list"
                    if style_key.startswith(("list", "список")) or _paragraph_has_numbering(child)
                    else "paragraph"
                )
                builder.add(kind, info.text, current_path())
            if info.has_image:
                block = builder.add("image", "[Изображение или диаграмма]", current_path(), metadata={"image": True})
                output.images += 1
                output.manual_review.append(
                    ManualReviewItem(
                        file=file,
                        section=block.section,
                        reason="Изображение или диаграмма не анализировались",
                        kind="image",
                        block_id=block.block_id,
                    )
                )
            if info.has_object:
                block = builder.add("image", "[Встроенный объект]", current_path(), metadata={"object": True})
                output.images += 1
                output.manual_review.append(
                    ManualReviewItem(
                        file=file,
                        section=block.section,
                        reason="Встроенный OLE-объект не анализировался",
                        kind="diagram",
                        block_id=block.block_id,
                    )
                )

        elif tag == "tbl":
            try:
                table = Table(child, document)
                markdown, rows, columns = _table_markdown(table)
            except Exception as exc:  # noqa: BLE001 - повреждённая таблица
                output.warnings.append(f"Не удалось разобрать таблицу в {file}: {exc}")
                output.manual_review.append(
                    ManualReviewItem(
                        file=file,
                        section=" / ".join(current_path()) or "(без раздела)",
                        reason=f"Таблица не разобрана: {exc}",
                        kind="unreadable",
                    )
                )
                continue
            if not markdown:
                continue
            builder.add(
                "table",
                markdown,
                current_path(),
                is_table=True,
                metadata={"rows": rows, "columns": columns},
            )
            output.tables += 1
            image_count = len(child.findall(".//" + qn("a:blip")))
            if image_count:
                output.images += image_count
                output.manual_review.append(
                    ManualReviewItem(
                        file=file,
                        section=" / ".join(current_path()) or "(без раздела)",
                        reason=f"В таблице найдено изображений: {image_count}; они не анализировались",
                        kind="image",
                    )
                )

    output.blocks = builder.blocks
    if not output.blocks:
        output.warnings.append(f"В {file} не найдено текстового содержимого")
    return output

